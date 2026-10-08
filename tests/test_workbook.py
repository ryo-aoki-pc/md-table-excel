"""ブックの書き出し・読み込み・取り込みの判定のテスト。"""

import datetime
import os
import re
import shutil
import tempfile
import unittest
import zipfile
from unittest import mock
from xml.sax.saxutils import escape

import helpers

import openpyxl

from mdtable import sync, workbook
from mdtable.model import normalize_text


class TempDir(object):
    def __enter__(self):
        self.path = tempfile.mkdtemp(prefix="mdtable-test-")
        return self.path

    def __exit__(self, *args):
        shutil.rmtree(self.path, ignore_errors=True)


def write(path, text):
    with open(path, "w", encoding="utf-8", newline="") as handle:
        handle.write(text)


def read(path):
    with open(path, encoding="utf-8", newline="") as handle:
        return handle.read()


def edit_sheet(xlsx, title, edits):
    """openpyxl でセルを書き換える（Excel で入力したのと同じく文字列として）。"""
    book = openpyxl.load_workbook(xlsx)
    sheet = book[title]
    for ref, value in edits.items():
        sheet[ref].value = value
        if isinstance(value, str):
            sheet[ref].data_type = "s"
    book.save(xlsx)


def excel_shape(xlsx):
    """openpyxl の書いたブックを、Excel が保存したときの形（共有文字列）に書き換える。

    Excel は CR を _x000D_ に、文字の _xHHHH_ を _x005F_xHHHH_ にして共有文字列に入れる。
    """
    with zipfile.ZipFile(xlsx) as source:
        entries = [(info, source.read(info.filename)) for info in source.infolist()]
    strings = []

    def to_shared(match):
        attrs, text = match.group(1), match.group(2)
        strings.append(text)
        attrs = attrs.replace(' t="inlineStr"', ' t="s"')
        return "<c%s><v>%d</v></c>" % (attrs, len(strings) - 1)

    pattern = re.compile(r'<c([^>]*t="inlineStr"[^>]*)><is><t(?: [^>]*)?>([\s\S]*?)</t></is></c>')
    out = []
    for info, data in entries:
        if info.filename.startswith("xl/worksheets/sheet"):
            data = pattern.sub(to_shared, data.decode("utf-8")).encode("utf-8")
        out.append((info, data))
    items = []
    for text in strings:
        # Windows の openpyxl はセル内の改行を XML に CRLF で書くので、先に LF にそろえる
        text = text.replace("\r\n", "\n")
        text = text.replace("_x000D_", "_x005F_x000D_").replace("\n", "_x000D_\n")
        items.append('<si><t xml:space="preserve">%s</t><rPh sb="0" eb="1"><t>ふりがな</t></rPh></si>' % text)
    shared = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
              '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="%d" '
              'uniqueCount="%d">%s</sst>' % (len(strings), len(strings), "".join(items)))
    with zipfile.ZipFile(xlsx, "w", zipfile.ZIP_DEFLATED) as target:
        for info, data in out:
            if info.filename == "[Content_Types].xml":
                data = data.decode("utf-8").replace(
                    "</Types>",
                    '<Override PartName="/xl/sharedStrings.xml" ContentType="application/'
                    'vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>'
                ).encode("utf-8")
            if info.filename == "xl/_rels/workbook.xml.rels":
                data = data.decode("utf-8").replace(
                    "</Relationships>",
                    '<Relationship Id="rIdShared" Type="http://schemas.openxmlformats.org/'
                    'officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/>'
                    "</Relationships>").encode("utf-8")
            target.writestr(info, data)
        target.writestr("xl/sharedStrings.xml", shared)


SAMPLE = """# 文書

| ツール | 説明 |
|---|---|
| lazygit | `scoop install lazygit` |
| yazi | 設定<br>- `yazi.toml` |

<table>
<tr><th>項目</th><th>値</th></tr>
<tr><td>版</td><td>3.4</td></tr>
</table>
"""


class NormalizeTest(unittest.TestCase):
    def test_normalize(self):
        self.assertEqual(normalize_text("a\r\nb\rc  \n\n"), "a\nb\nc")
        self.assertEqual(normalize_text("\n\n  - a"), "  - a")
        self.assertEqual(normalize_text("a_x000D_\nb"), "a\nb")

    def test_escape_round_trip(self):
        # W1: 文字の _x000D_ は _x005F_ を付けて書き、読み戻すと元に戻る
        for text in ("_x000D_", "a_x0041_b", "plain"):
            self.assertEqual(workbook.unescape_from_excel(workbook.escape_for_excel(text)), text)

    def test_cell_checks(self):
        # W2
        self.assertIsNotNone(workbook.check_cell_text("a\x0cb"))
        self.assertIsNotNone(workbook.check_cell_text("x" * 32768))
        self.assertIsNone(workbook.check_cell_text("x" * 32767))

    def test_sheet_names(self):
        used = set()
        name = workbook.sheet_name_for(3, "見出し: [とても] 長い/見出し*の名前" * 3, used)
        self.assertLessEqual(len(name.encode("utf-16-le")) // 2, 31)
        self.assertNotRegex(name, r"[\\/?*\[\]:]")
        self.assertNotEqual(workbook.sheet_name_for(3, "見出し: [とても] 長い/見出し*の名前" * 3, used), name)


class FontTest(unittest.TestCase):
    def test_font_entry_matches(self):
        family = "HackGen Console NF"
        for entry in ("HackGen Console NF Regular (TrueType)", "HackGen Console NF Bold (TrueType)",
                      "Foo & HackGen Console NF (TrueType)", "hackgen console nf"):
            self.assertTrue(workbook.font_entry_matches(entry, family), entry)
        for entry in ("HackGen35 Console NF Regular (TrueType)", "HackGen NF Regular (TrueType)",
                      "HackGen Console NFX Regular (TrueType)", "BIZ UDGothic & BIZ UDPGothic (TrueType)"):
            self.assertFalse(workbook.font_entry_matches(entry, family), entry)

    def test_font_name(self):
        # 環境変数が最優先。無ければ HackGen Console NF が入っているかで決める
        with mock.patch.dict(os.environ, {"MDTABLE_FONT": "Consolas"}):
            with mock.patch.object(workbook, "font_installed", return_value=True):
                self.assertEqual(workbook.font_name(), "Consolas")
        with mock.patch.dict(os.environ):
            os.environ.pop("MDTABLE_FONT", None)
            with mock.patch.object(workbook, "font_installed", return_value=True):
                self.assertEqual(workbook.font_name(), "HackGen Console NF")
            with mock.patch.object(workbook, "font_installed", return_value=False):
                self.assertEqual(workbook.font_name(), "BIZ UDゴシック")

    @unittest.skipUnless(os.name == "nt", "Windows のフォントの登録を見る")
    def test_registered_fonts(self):
        self.assertFalse(workbook.font_installed("mdtable に無い書体 0123"))
        self.assertIsInstance(workbook.font_installed(workbook.PREFERRED_FONT), bool)


class ExportImportTest(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="mdtable-test-")
        self.md = os.path.join(self.dir, "文書 1.md")
        self.xlsx = os.path.join(self.dir, "文書 1.tables.xlsx")
        write(self.md, SAMPLE)

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def export(self, **kwargs):
        return sync.export(self.md, self.xlsx, force=True, **kwargs)

    def sheet(self, number):
        book = openpyxl.load_workbook(self.xlsx)
        return [t for t in book.sheetnames if t.startswith("表%d" % number)][0]

    def test_no_edit_round_trip(self):
        self.export()
        report = sync.import_workbook(self.xlsx)
        self.assertFalse(report.changed)
        self.assertEqual(read(self.md), SAMPLE)

    def test_excel_shaped_round_trip(self):
        # Excel が保存した形（共有文字列・_x000D_・ふりがな）でも変わらない
        self.export()
        excel_shape(self.xlsx)
        report = sync.import_workbook(self.xlsx)
        self.assertEqual(report.errors, [])
        self.assertFalse(report.changed)

    def test_text_cells(self):
        self.export()
        book = openpyxl.load_workbook(self.xlsx)
        sheet = book[self.sheet(1)]
        self.assertEqual(sheet["B3"].value, "設定\n\\- `yazi.toml`")
        self.assertEqual(sheet["B3"].number_format, "@")
        self.assertTrue(sheet["A1"].font.b)
        self.assertEqual(book[workbook.META_SHEET].sheet_state, "veryHidden")

    def test_font(self):
        with mock.patch.dict(os.environ, {"MDTABLE_FONT": "Consolas"}):
            self.export()
        book = openpyxl.load_workbook(self.xlsx)
        sheet = book[self.sheet(1)]
        index = book[workbook.INDEX_SHEET]
        for cell in (sheet["A1"], sheet["B2"], index["A1"], index["B4"]):
            self.assertEqual(cell.font.name, "Consolas", cell.coordinate)
        self.assertEqual(sheet.column_dimensions["A"].font.name, "Consolas")

    def test_edit_and_import(self):
        self.export()
        edit_sheet(self.xlsx, self.sheet(1), {"B2": "`scoop install lazygit`\n- 確認する"})
        report = sync.import_workbook(self.xlsx)
        self.assertEqual(report.errors, [])
        self.assertIn("| lazygit | `scoop install lazygit`<ul><li>確認する</li></ul> |", read(self.md))

    def test_numbers_and_dates(self):
        # W4: 標準の書式の数値は文字にして警告、日付は止める（--force で書き戻す）
        self.export()
        book = openpyxl.load_workbook(self.xlsx)
        sheet = book[self.sheet(2)]
        sheet["B2"].value = 3.5
        sheet["B2"].number_format = "General"
        book.save(self.xlsx)
        report = sync.import_workbook(self.xlsx)
        self.assertEqual(report.errors, [])
        self.assertTrue(any("数値" in w for w in report.warnings))
        self.assertIn("<td>3.5</td>", read(self.md))

        self.export()
        book = openpyxl.load_workbook(self.xlsx)
        sheet = book[self.sheet(2)]
        sheet["B2"].value = datetime.datetime(2026, 10, 9)
        book.save(self.xlsx)
        before = read(self.md)
        report = sync.import_workbook(self.xlsx)
        self.assertTrue(report.errors)
        self.assertEqual(read(self.md), before)
        report = sync.import_workbook(self.xlsx, force=True)
        self.assertIn("<td>2026-10-09</td>", read(self.md))

    def test_formula_like_text(self):
        self.export()
        edit_sheet(self.xlsx, self.sheet(2), {"B2": "=A1"})
        book = openpyxl.load_workbook(self.xlsx)
        self.assertEqual(book[self.sheet(2)]["B2"].data_type, "s")
        report = sync.import_workbook(self.xlsx)
        self.assertEqual(report.errors, [])
        self.assertIn("<td>=A1</td>", read(self.md))

    def test_renamed_and_deleted_sheets(self):
        self.export()
        book = openpyxl.load_workbook(self.xlsx)
        old = self.sheet(1)
        sheet = book[old]
        sheet.title = "名前を変えた"
        # Excel はシート名を変えると名前の定義も直す（openpyxl は直さないので合わせる）
        book.defined_names["mdtable_1"].attr_text = "'名前を変えた'!$A$1:$B$3"
        sheet["A2"].value = "lazygit2"
        sheet["A2"].data_type = "s"
        del book[self.sheet(2)]
        book.save(self.xlsx)
        report = sync.import_workbook(self.xlsx)
        self.assertIn("| lazygit2 |", read(self.md))
        self.assertTrue(any("シートが見つからない" in w for w in report.warnings))

    def test_far_values_are_ignored(self):
        self.export()
        edit_sheet(self.xlsx, self.sheet(1), {"Z100": "stray"})
        report = sync.import_workbook(self.xlsx)
        self.assertFalse(report.changed)
        self.assertTrue(any("離れた位置" in w for w in report.warnings))

    def test_adjacent_row_extends(self):
        self.export()
        edit_sheet(self.xlsx, self.sheet(1), {"A4": "delta", "B4": "使う"})
        sync.import_workbook(self.xlsx)
        self.assertIn("| delta | 使う |", read(self.md))

    def test_conflict_and_force(self):
        # C3: 書き出した後に Markdown の表を変えると止まる
        self.export()
        write(self.md, SAMPLE.replace("| yazi |", "| yazi2 |"))
        edit_sheet(self.xlsx, self.sheet(1), {"A2": "lazygit3"})
        report = sync.import_workbook(self.xlsx)
        self.assertTrue(any("衝突" in e for e in report.errors))
        self.assertNotIn("lazygit3", read(self.md))
        report = sync.import_workbook(self.xlsx, force=True)
        self.assertIn("lazygit3", read(self.md))

    def test_applied_twice(self):
        # C2: 同じブックをもう一度取り込むと「適用済み」
        self.export()
        edit_sheet(self.xlsx, self.sheet(1), {"A2": "lazygit4"})
        sync.import_workbook(self.xlsx)
        report = sync.import_workbook(self.xlsx)
        self.assertEqual(report.errors, [])
        self.assertFalse(report.changed)
        self.assertTrue(any("適用済み" in m for m in report.messages))

    def test_applied_twice_after_normalizing(self):
        # 表の下に足した行のセルの配置が列と違う（多いほうにそろえる）、セルの文字が収束する形に
        # 直る（- y⏎z → - y⏎  z）のどちらがあっても、もう一度取り込むと「適用済み」
        write(self.md, "| a | b |\n|---|:---:|\n| 1 | 2 |\n")
        self.export()
        edit_sheet(self.xlsx, self.sheet(1), {"A3": "x", "B3": "- y\nz"})
        report = sync.import_workbook(self.xlsx)
        self.assertEqual(report.errors, [])
        self.assertTrue(any("配置が混在" in w for w in report.warnings))
        self.assertIn("| x | <ul><li>y<br>z</li></ul> |", read(self.md))
        after = read(self.md)
        report = sync.import_workbook(self.xlsx)
        self.assertEqual(report.errors, [])
        self.assertFalse(report.changed)
        self.assertTrue(any("適用済み" in m for m in report.messages))
        self.assertEqual(read(self.md), after)

    def test_many_tables_with_body(self):
        # 本文と 7 つの表がある見本（docs/verification/multi-tables.md）で 5 つの表を編集しても、
        # 表の外の行（足した空行を除く）と、編集していない表は 1 文字も変わらない
        with open(os.path.join(helpers.ROOT, "docs", "verification", "multi-tables.md"),
                  encoding="utf-8", newline="") as handle:
            original = handle.read()
        write(self.md, original)
        self.export()
        edit_sheet(self.xlsx, self.sheet(1), {"B3": "0.2.0", "A4": "フォント", "B4": "HackGen Console NF",
                                              "C4": "- 入っていれば使う\n- 無ければ BIZ UDゴシック"})
        edit_sheet(self.xlsx, self.sheet(2), {"B2": "`python scripts/setup-mdtable.py`\n（リポジトリの外に入れる）"})
        edit_sheet(self.xlsx, self.sheet(3), {"B3": "セル結合ができる\n\n```html\n<td rowspan=\"2\">結合</td>\n```"})
        edit_sheet(self.xlsx, self.sheet(4), {"C3": "```sh\nsudo dnf install -y epel-release\n```",
                                              "C4": "ln\n: シンボリックリンクを作る"})
        edit_sheet(self.xlsx, self.sheet(5), {"B2": "編集（**Excel** で開く）"})
        report = sync.import_workbook(self.xlsx)
        self.assertEqual(report.errors, [])
        edited = read(self.md)

        def split(text):
            regions = helpers.regions(text)
            lines = text.split("\n")
            covered = {i for r in regions for i in range(r.start, r.end + 1)}
            outside = [line for i, line in enumerate(lines) if i not in covered and line.strip()]
            tables = ["\n".join(lines[r.start:r.end + 1]) for r in regions]
            return outside, tables, [r.kind for r in regions]

        outside_before, tables_before, _ = split(original)
        outside_after, tables_after, kinds = split(edited)
        self.assertEqual(outside_after, outside_before)
        self.assertEqual(len(tables_after), 7)
        self.assertEqual([i + 1 for i in range(7) if tables_before[i] != tables_after[i]], [1, 2, 3, 4, 5])
        self.assertEqual(kinds, ["pipe", "pipe", "html", "html", "pipe", "pipe", "pipe"])
        # HTML にした表 3 の直後のリストは、空行を足してリストのまま
        self.assertIn("</table>\n\n- 表のすぐ後のリスト（空行なし）", edited)
        report = sync.import_workbook(self.xlsx)
        self.assertFalse(report.changed)
        self.assertEqual(len([m for m in report.messages if "適用済み" in m]), 5)

    def test_merged_cells_keep_text_format(self):
        # 結合で隠れるセルも、左上のセルと同じ書式（文字列の書式）にする。Excel で結合を解いて
        # 入力しても文字列のまま。openpyxl は読むときに隠れるセルの書式を捨てるので、XML で確かめる
        write(self.md, '<table>\n<tr><th rowspan="2">a</th><th>b</th></tr>\n<tr><td>c</td></tr>\n</table>\n')
        self.export()
        with zipfile.ZipFile(self.xlsx) as book:
            sheet = book.read("xl/worksheets/sheet2.xml").decode("utf-8")
            styles = book.read("xl/styles.xml").decode("utf-8")
        self.assertIn('<mergeCell ref="A1:A2"', sheet)
        top = re.search(r'<c r="A1" s="(\d+)"', sheet).group(1)
        covered = re.search(r'<c r="A2" s="(\d+)"', sheet).group(1)
        self.assertEqual(covered, top)
        xfs = re.findall(r"<xf [^>]*>", re.search(r"<cellXfs[^>]*>(.*?)</cellXfs>", styles, re.S).group(1))
        self.assertIn('numFmtId="49"', xfs[int(covered)])
        report = sync.import_workbook(self.xlsx)
        self.assertEqual(report.errors, [])
        self.assertFalse(report.changed)

    def test_column_style_alignment(self):
        # パイプ表の列の配置を列のスタイルにも付け、表の下の行に入力したセルも同じ配置にする
        write(self.md, "| a | b |\n|---|:---:|\n| 1 | 2 |\n")
        self.export()
        sheet = openpyxl.load_workbook(self.xlsx)[self.sheet(1)]
        self.assertEqual(sheet.column_dimensions["B"].alignment.horizontal, "center")
        self.assertIsNone(sheet.column_dimensions["A"].alignment.horizontal)

    def test_line_endings_and_bom(self):
        # C4・D17: 改行コードと BOM を保ち、改行コードを変えても衝突しない
        text = SAMPLE.replace("\n", "\r\n")
        with open(self.md, "wb") as handle:
            handle.write(b"\xef\xbb\xbf" + text.encode("utf-8"))
        self.export()
        edit_sheet(self.xlsx, self.sheet(1), {"A2": "lazygit5"})
        sync.import_workbook(self.xlsx)
        with open(self.md, "rb") as handle:
            data = handle.read()
        self.assertTrue(data.startswith(b"\xef\xbb\xbf"))
        self.assertNotIn(b"\n", data.replace(b"\r\n", b""))
        self.assertIn("lazygit5".encode("utf-8"), data)

    def test_identical_tables(self):
        # C1: 同じ表が 2 つあっても、編集したほうだけが変わる
        table = "| a | b |\n|---|---|\n| 1 | 2 |\n"
        write(self.md, table + "\n" + table)
        self.export()
        edit_sheet(self.xlsx, self.sheet(2), {"A2": "x"})
        sync.import_workbook(self.xlsx)
        self.assertEqual(read(self.md), table + "\n" + table.replace("| 1 |", "| x |"))

    def test_conversion_then_reimport(self):
        # C5: パイプ表を HTML の表にした後、同じブックをもう一度取り込んでも変わらない
        self.export()
        book = openpyxl.load_workbook(self.xlsx)
        book[self.sheet(1)].merge_cells("A2:A3")
        book.save(self.xlsx)
        report = sync.import_workbook(self.xlsx)
        self.assertEqual(report.errors, [])
        self.assertIn('<td rowspan="2">lazygit</td>', read(self.md))
        after = read(self.md)
        report = sync.import_workbook(self.xlsx)
        self.assertFalse(report.changed)
        self.assertEqual(read(self.md), after)

    def test_control_character_is_excluded(self):
        # W2: 書き出せない表は対象外にして、ほかの表は書き出す
        write(self.md, "| a |\n|---|\n| x\x0cy |\n\n| b |\n|---|\n| 1 |\n")
        report = self.export()
        self.assertTrue(any("対象外" in w for w in report.warnings))
        book = openpyxl.load_workbook(self.xlsx)
        self.assertEqual(len([t for t in book.sheetnames if t.startswith("表")]), 1)


if __name__ == "__main__":
    unittest.main()
