"""Excel のブックの書き出しと読み込み。

ブックの形:
- 「目次」シート: 表の一覧（シートへのリンク）、対象外の表、セルの書き方の早見表
- 表ごとのシート: 表は A1 から。セルは文字列（@）。見出しセルは太字 + 塗りつぶし、結合は rowspan / colspan
- 隠しシート _mdtable: 元の Markdown のパスと、表ごとの位置・ハッシュ・指紋
- 名前の定義 mdtable_<番号>: 表の範囲（シート名の変更に追従する）
"""

import datetime
import os
import re
import shutil
import subprocess
import tempfile
import warnings

from .model import Cell, TableModel, normalize_text
from .textwidth import text_width, wrapped_line_count

FORMAT_VERSION = 1
META_SHEET = "_mdtable"
INDEX_SHEET = "目次"
# セルのフォント。HackGen Console NF（日本語の等幅で、\ が ¥ に見えない）が入っていればそれ、
# 無ければどの Windows にもある BIZ UDゴシック（\ は ¥ に見える）。環境変数 MDTABLE_FONT で変えられる
PREFERRED_FONT = "HackGen Console NF"
FALLBACK_FONT = "BIZ UDゴシック"
FONT_SIZE = 10.5
HEADER_FILL = "D9D9D9"
BORDER_COLOR = "A6A6A6"
EXCEL_TEXT_LIMIT = 32767
NAME_PREFIX = "mdtable_"

_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_ESCAPE_LIKE = re.compile(r"_(?=x[0-9A-Fa-f]{4}_)")
_ESCAPED = re.compile(r"_x005F_(?=x[0-9A-Fa-f]{4}_)", re.I)
_SHEET_FORBIDDEN = re.compile(r"[\\/?*\[\]:]")

QUICK_REFERENCE = [
    "セルの書き方（GLFM の Markdown）",
    "・セル内の改行（Alt+Enter）は表示上の改行。空行は段落の区切り",
    "・箇条書き「- 」、番号付き「1. 」、入れ子は親の本文の位置まで空白で字下げ。リストは空行で終わる",
    "・説明リスト: 「用語」の次の行に「: 説明」。組と組の間は空行",
    "・コードブロック: ``` で囲む（言語名も書ける）",
    "・「|」はそのまま書く。行頭の「-」「1.」「>」「:」などを文字として書くときは「\\-」のようにエスケープする",
    "・「x」の次の行に「2. y」と書くと 1 つの段落になる（番号付きリストは 1 から始めるか、前に空行を入れる）",
    "・Markdown の表: 1 行目が見出し。列の配置は列の中で多いものになる",
    "・HTML の表: 太字かつ塗りつぶしのセルが見出し（<th>）。セル結合は rowspan / colspan",
    "・セル結合やコードブロックがある Markdown の表は、書き戻すと HTML の表になる",
    "・行を挿入すると上の行の書式が写る。見出しの直下に挿入したら [挿入オプション] で「下と同じ書式を適用」を選ぶ",
]


class WorkbookError(Exception):
    pass


class ExportTable(object):
    """書き出す 1 つの表。"""

    def __init__(self, region, model, occurrence):
        self.region = region
        self.model = model
        self.occurrence = occurrence
        self.sheet_name = None


class SheetTable(object):
    """ブックから読んだ 1 つの表。"""

    def __init__(self, number):
        self.number = number
        self.kind = None
        self.start = None
        self.end = None
        self.source_hash = None
        self.fingerprint = None
        self.sheet_title = None
        self.occurrence = 0
        self.model = None
        self.missing = None      # シートが見つからない理由
        self.errors = []         # 取り込みを止める問題
        self.warnings = []


class WorkbookData(object):
    def __init__(self):
        self.markdown_path = None
        self.markdown_relpath = None
        self.format_version = None
        self.tool_version = None
        self.tables = []
        self.warnings = []


def _openpyxl():
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        raise WorkbookError("openpyxl がありません。導入手順（README）を確かめてください")
    import openpyxl
    return openpyxl


# --- フォント ---

_FONT_KEY = r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts"
_FONT_KIND = re.compile(r"\s*\([^()]*\)\s*$")
_installed_cache = {}


def font_entry_matches(entry, family):
    """Windows のフォントの登録名（例: HackGen Console NF Bold (TrueType)）が family の書体か。"""
    family = family.lower()
    for name in _FONT_KIND.sub("", entry).split("&"):
        name = name.strip().lower()
        if name == family or name.startswith(family + " "):
            return True
    return False


def _registered_fonts():
    """Windows に登録されたフォントの名前（全ユーザー向けと、自分だけに入れたもの）。"""
    import winreg
    names = []
    for root in (winreg.HKEY_LOCAL_MACHINE, winreg.HKEY_CURRENT_USER):
        try:
            key = winreg.OpenKey(root, _FONT_KEY)
        except OSError:
            continue
        with key:
            index = 0
            while True:
                try:
                    names.append(winreg.EnumValue(key, index)[0])
                except OSError:
                    break
                index += 1
    return names


def _fontconfig_families():
    command = shutil.which("fc-list")
    if not command:
        return []
    try:
        result = subprocess.run([command, ":", "family"], capture_output=True, timeout=10)
    except (OSError, subprocess.SubprocessError):
        return []
    families = []
    for line in result.stdout.decode("utf-8", "replace").splitlines():
        families.extend(part.strip() for part in line.split(","))
    return families


def font_installed(family):
    """この PC にそのフォントが入っているか（結果は覚えておく）。"""
    if family not in _installed_cache:
        if os.name == "nt":
            found = any(font_entry_matches(entry, family) for entry in _registered_fonts())
        else:
            found = family.lower() in [name.lower() for name in _fontconfig_families()]
        _installed_cache[family] = found
    return _installed_cache[family]


def font_name():
    """ブックのフォント。MDTABLE_FONT → HackGen Console NF（入っていれば）→ BIZ UDゴシック。"""
    name = os.environ.get("MDTABLE_FONT")
    if name:
        return name
    return PREFERRED_FONT if font_installed(PREFERRED_FONT) else FALLBACK_FONT


def check_cell_text(text):
    """Excel のセルに書けるか。書けなければ理由を返す。"""
    if _ILLEGAL.search(text):
        return "制御文字が入っている"
    if len(text.encode("utf-16-le")) // 2 > EXCEL_TEXT_LIMIT:
        return "Excel の 1 セルの上限（32,767 文字）を超える"
    return None


def escape_for_excel(text):
    """_xHHHH_ の形の文字を、Excel が文字に戻さないようにエスケープする。"""
    return _ESCAPE_LIKE.sub("_x005F_", text)


def unescape_from_excel(text):
    return _ESCAPED.sub("_", text)


def sheet_name_for(number, heading, used):
    base = "表%d" % number
    title = heading.strip() if heading else ""
    title = _SHEET_FORBIDDEN.sub("", title)
    title = re.sub(r"\s+", " ", title).strip(" '")
    name = base + (" " + title if title else "")

    def units(text):
        return len(text.encode("utf-16-le")) // 2

    while units(name) > 31:
        name = name[:-1]
    name = name.rstrip(" '")
    candidate = name
    counter = 2
    while candidate.lower() in used or candidate.lower() == "history":
        suffix = "(%d)" % counter
        trimmed = name
        while units(trimmed + suffix) > 31:
            trimmed = trimmed[:-1]
        candidate = trimmed + suffix
        counter += 1
    used.add(candidate.lower())
    return candidate


def quote_sheet(name):
    return "'" + name.replace("'", "''") + "'"


def _header_rows(model):
    count = 0
    for row in range(model.rows):
        anchors = [c for c in model.cells if c.row == row]
        if anchors and all(c.header for c in anchors):
            count += 1
        else:
            break
    return count


def _column_widths(model):
    widths = []
    for col in range(model.cols):
        width = 0
        for cell in model.cells:
            if cell.col == col and cell.colspan == 1:
                for line in cell.text.split("\n"):
                    width = max(width, text_width(line.expandtabs(4)))
        widths.append(min(60, max(6, width + 2)) if width else 8)
    return widths


def _column_alignments(model):
    """列の 1 列分のセルがすべて同じ配置なら、その配置（無ければ None）。"""
    result = []
    for col in range(model.cols):
        aligns = {cell.align for cell in model.cells if cell.col == col and cell.colspan == 1}
        result.append(aligns.pop() if len(aligns) == 1 else None)
    return result


def _write_table_sheet(openpyxl, ws, model, family):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    side = Side(style="thin", color=BORDER_COLOR)
    border = Border(left=side, right=side, top=side, bottom=side)
    body_font = Font(name=family, size=FONT_SIZE)
    header_font = Font(name=family, size=FONT_SIZE, bold=True)
    header_fill = PatternFill(fill_type="solid", fgColor=HEADER_FILL)
    widths = _column_widths(model)
    aligns = _column_alignments(model)

    for index in range(model.cols + 5):
        letter = get_column_letter(index + 1)
        dimension = ws.column_dimensions[letter]
        dimension.width = widths[index] if index < len(widths) else 8
        dimension.number_format = "@"
        dimension.font = body_font
        # 表の下の行に入力したセルも、列の配置になるようにする
        dimension.alignment = Alignment(wrap_text=True, vertical="top",
                                        horizontal=aligns[index] if index < len(aligns) else None)

    # 結合を先にする。結合で隠れるセルにも下で書式を付け、結合を解いても文字列の書式のままにする
    for cell in model.cells:
        if cell.rowspan > 1 or cell.colspan > 1:
            ws.merge_cells(start_row=cell.row + 1, start_column=cell.col + 1,
                           end_row=cell.row + cell.rowspan, end_column=cell.col + cell.colspan)

    for row in range(model.rows):
        for col in range(model.cols):
            target = ws.cell(row=row + 1, column=col + 1)
            target.number_format = "@"
            target.border = border
            target.font = body_font
            target.alignment = Alignment(wrap_text=True, vertical="top")

    for cell in model.cells:
        target = ws.cell(row=cell.row + 1, column=cell.col + 1)
        if cell.text:
            target.value = escape_for_excel(cell.text)
            target.data_type = "s"
        font = header_font if cell.header else body_font
        alignment = Alignment(wrap_text=True, vertical="top", horizontal=cell.align)
        for r in range(cell.row, cell.row + cell.rowspan):
            for c in range(cell.col, cell.col + cell.colspan):
                part = ws.cell(row=r + 1, column=c + 1)
                part.font = font
                part.alignment = alignment
                if cell.header:
                    part.fill = header_fill

    # 結合を含む行は Excel が高さを自動で合わせないので見積もる
    heights = {}
    for cell in model.cells:
        if cell.rowspan == 1 and cell.colspan == 1:
            continue
        width = sum(widths[c] for c in range(cell.col, cell.col + cell.colspan)) - 1
        lines = wrapped_line_count(cell.text, int(width)) if cell.text else 1
        per_row = -(-lines // cell.rowspan)
        for r in range(cell.row, cell.row + cell.rowspan):
            heights[r] = max(heights.get(r, 1), per_row)
    for cell in model.cells:
        if cell.row in heights and cell.rowspan == 1 and cell.colspan == 1:
            lines = wrapped_line_count(cell.text, int(widths[cell.col] - 1)) if cell.text else 1
            heights[cell.row] = max(heights[cell.row], lines)
    for row, lines in heights.items():
        ws.row_dimensions[row + 1].height = min(409, max(15, 15 * lines))

    header_rows = _header_rows(model)
    if 0 < header_rows < model.rows:
        ws.freeze_panes = "A%d" % (header_rows + 1)


def _write_index(openpyxl, ws, md_path, tables, excluded, family):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.hyperlink import Hyperlink

    bold = Font(name=family, size=FONT_SIZE, bold=True)
    normal = Font(name=family, size=FONT_SIZE)
    link_font = Font(name=family, size=FONT_SIZE, color="0563C1", underline="single")
    fill = PatternFill(fill_type="solid", fgColor=HEADER_FILL)

    ws["A1"] = "Markdown"
    ws["A1"].font = bold
    ws["B1"] = md_path
    ws["B1"].font = normal
    headers = ["表", "シート", "種類", "行", "場所", "大きさ", "直前の見出し"]
    for col, title in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=title)
        cell.font = bold
        cell.fill = fill
    row = 4
    for table in tables:
        region = table.region
        values = [
            region.number, table.sheet_name,
            "Markdown の表" if region.kind == "pipe" else "HTML の表",
            "%d〜%d" % (region.start + 1, region.end + 1),
            region.location or "",
            "%d 行 × %d 列" % (table.model.rows, table.model.cols),
            region.heading or "",
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.font = normal
            if isinstance(value, str):
                cell.data_type = "s"
        link = ws.cell(row=row, column=2)
        link.hyperlink = Hyperlink(ref=link.coordinate, location=NAME_PREFIX + str(region.number))
        link.font = link_font
        row += 1
    if excluded:
        row += 1
        ws.cell(row=row, column=1, value="対象外の表").font = bold
        row += 1
        for region in excluded:
            ws.cell(row=row, column=1, value=region.number).font = normal
            ws.cell(row=row, column=4, value="%d〜%d" % (region.start + 1, region.end + 1)).font = normal
            reason = ws.cell(row=row, column=5, value=region.excluded)
            reason.font = normal
            reason.data_type = "s"
            row += 1
    row += 1
    for index, line in enumerate(QUICK_REFERENCE):
        cell = ws.cell(row=row, column=1, value=line)
        cell.font = bold if index == 0 else normal
        cell.data_type = "s"
        row += 1
    # 1 列目は表の番号だが、A1 の「Markdown」が切れない幅にする
    widths = [10, 28, 14, 12, 30, 14, 40]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width
    for line in ws.iter_rows(min_row=4, max_row=row):
        for cell in line:
            cell.alignment = Alignment(vertical="top")


def _write_meta(ws, md_path, xlsx_path, tables, tool_version):
    rel = None
    try:
        rel = os.path.relpath(md_path, os.path.dirname(os.path.abspath(xlsx_path)))
    except ValueError:
        rel = None
    rows = [
        ["md-table-excel", FORMAT_VERSION, tool_version],
        ["markdown", md_path, rel.replace("\\", "/") if rel else ""],
        ["exported", datetime.datetime.now().replace(microsecond=0).isoformat()],
        [],
        ["number", "kind", "start", "end", "source_sha256", "fingerprint", "sheet", "occurrence"],
    ]
    for table in tables:
        region = table.region
        rows.append([region.number, region.kind, region.start + 1, region.end + 1,
                     region.source_hash(), table.model.fingerprint(), table.sheet_name,
                     table.occurrence])
    for r, values in enumerate(rows, 1):
        for c, value in enumerate(values, 1):
            cell = ws.cell(row=r, column=c, value=value)
            if isinstance(value, str):
                cell.data_type = "s"


def write(path, md_path, tables, excluded, tool_version):
    """ブックを書く（同じフォルダーの一時ファイルに書いてから置き換える）。"""
    openpyxl = _openpyxl()
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.utils import get_column_letter

    family = font_name()
    workbook = openpyxl.Workbook()
    index_sheet = workbook.active
    index_sheet.title = INDEX_SHEET
    used = {INDEX_SHEET.lower(), META_SHEET.lower()}
    for table in tables:
        table.sheet_name = sheet_name_for(table.region.number, table.region.heading, used)
        ws = workbook.create_sheet(table.sheet_name)
        _write_table_sheet(openpyxl, ws, table.model, family)
        ref = "%s!$A$1:$%s$%d" % (quote_sheet(table.sheet_name),
                                   get_column_letter(max(1, table.model.cols)), max(1, table.model.rows))
        name = NAME_PREFIX + str(table.region.number)
        workbook.defined_names[name] = DefinedName(name, attr_text=ref)
    _write_index(openpyxl, index_sheet, md_path, tables, excluded, family)
    meta = workbook.create_sheet(META_SHEET)
    _write_meta(meta, md_path, path, tables, tool_version)
    meta.sheet_state = "veryHidden"

    target = os.path.abspath(path)
    directory = os.path.dirname(target)
    handle, temp_path = tempfile.mkstemp(prefix=".mdtable-", suffix=".xlsx", dir=directory)
    os.close(handle)
    try:
        workbook.save(temp_path)
        from .document import replace_with_retry
        try:
            replace_with_retry(temp_path, target)
        except PermissionError:
            raise WorkbookError("ブックを書けません（Excel で開いていませんか）: %s" % target)
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)


# --- 読み込み ---

def _format_number(value):
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if value.is_integer() and abs(value) < 1e15:
        return str(int(value))
    return repr(value)


def _format_datetime(value):
    if isinstance(value, datetime.datetime):
        if value.time() == datetime.time(0, 0):
            return value.date().isoformat()
        return value.replace(microsecond=0).isoformat(sep=" ")
    return value.isoformat()


def _cell_text(cell, table, force, where):
    from openpyxl.cell.rich_text import CellRichText

    value = cell.value
    if value is None:
        return ""
    if isinstance(value, CellRichText):
        table.warnings.append("%s: セルの一部だけの書式は書き戻さない（文字だけを使う）" % where)
        value = str(value)
    if isinstance(value, str):
        if cell.data_type == "f":
            table.warnings.append("%s: 数式は式の文字のまま書き戻す: %s" % (where, value))
        elif cell.data_type == "e":
            table.warnings.append("%s: エラー値は文字のまま書き戻す: %s" % (where, value))
        return unescape_from_excel(normalize_text(value))
    if isinstance(value, bool):
        table.warnings.append("%s: 真偽値を文字にした: %s" % (where, _format_number(value)))
        return _format_number(value)
    if isinstance(value, (int, float)):
        text = _format_number(value)
        if cell.number_format in ("General", "@"):
            table.warnings.append("%s: 数値を文字にした: %s（文字列の書式で入力し直すと確実）" % (where, text))
        elif force:
            table.warnings.append("%s: 数値を文字にした: %s" % (where, text))
        else:
            table.errors.append(
                "%s: Excel が数値に変えたセルがある（%s、書式 %s）。セルを文字列の書式にして入力し直すか、--force で数値のまま文字にする"
                % (where, text, cell.number_format))
        return text
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        text = _format_datetime(value)
        if force:
            table.warnings.append("%s: 日付・時刻を文字にした: %s" % (where, text))
        else:
            table.errors.append(
                "%s: Excel が日付・時刻に変えたセルがある（%s）。セルを文字列の書式にして入力し直すか、--force で %s として書き戻す"
                % (where, text, text))
        return text
    return normalize_text(str(value))


def _horizontal(cell):
    value = cell.alignment.horizontal if cell.alignment is not None else None
    if value in (None, "general", "fill"):
        return None
    if value in ("center", "centerContinuous"):
        return "center"
    if value in ("justify", "distributed"):
        return "justify"
    return value


def _is_header(cell):
    font = cell.font
    fill = cell.fill
    return bool(font is not None and font.b) and fill is not None and fill.fill_type == "solid"


def _parse_range(cells):
    from openpyxl.utils.cell import range_boundaries

    text = cells.replace("$", "")
    if not text or "#REF!" in text:
        return None
    try:
        min_col, min_row, max_col, max_row = range_boundaries(text)
    except ValueError:
        return None
    return min_row, min_col, max_row, max_col


def _table_extent(ws, start):
    """表の範囲（上・左・下・右、1 起点）と、無視した離れた値の位置。"""
    top, left, bottom, right = start
    values = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and str(cell.value) != "":
                values[(cell.row, cell.column)] = True
    merges = [(m.min_row, m.min_col, m.max_row, m.max_col) for m in ws.merged_cells.ranges]
    changed = True
    while changed:
        changed = False
        for r1, c1, r2, c2 in merges:
            if r1 <= bottom + 1 and r2 >= top and c1 <= right + 1 and c2 >= left:
                if r1 < top or c1 < left or r2 > bottom or c2 > right:
                    top, left = min(top, r1), min(left, c1)
                    bottom, right = max(bottom, r2), max(right, c2)
                    changed = True
        if any(values.get((bottom + 1, c)) for c in range(left, right + 1)):
            bottom += 1
            changed = True
        if any(values.get((r, right + 1)) for r in range(top, bottom + 1)):
            right += 1
            changed = True
    ignored = sorted(pos for pos in values if not (top <= pos[0] <= bottom and left <= pos[1] <= right))
    return (top, left, bottom, right), ignored


def _read_table_sheet(ws, start, table, force):
    from openpyxl.utils import get_column_letter

    (top, left, bottom, right), ignored = _table_extent(ws, start)
    if ignored:
        names = ", ".join("%s%d" % (get_column_letter(c), r) for r, c in ignored[:5])
        table.warnings.append("シート「%s」: 表から離れた位置の値を無視した（%s）" % (ws.title, names))
    merges = {}
    covered = set()
    for m in ws.merged_cells.ranges:
        if m.min_row < top or m.min_col < left or m.max_row > bottom or m.max_col > right:
            continue
        merges[(m.min_row, m.min_col)] = (m.max_row - m.min_row + 1, m.max_col - m.min_col + 1)
        for r in range(m.min_row, m.max_row + 1):
            for c in range(m.min_col, m.max_col + 1):
                if (r, c) != (m.min_row, m.min_col):
                    covered.add((r, c))
    cells = []
    for r in range(top, bottom + 1):
        for c in range(left, right + 1):
            if (r, c) in covered:
                continue
            source = ws.cell(row=r, column=c)
            rowspan, colspan = merges.get((r, c), (1, 1))
            where = "シート「%s」の %s%d" % (ws.title, get_column_letter(c), r)
            text = _cell_text(source, table, force, where)
            cells.append(Cell(r - top, c - left, text, rowspan, colspan,
                              header=_is_header(source), align=_horizontal(source)))
    return TableModel(None, bottom - top + 1, right - left + 1, cells)


def _read_meta(ws, data):
    rows = [[cell.value for cell in row] for row in ws.iter_rows()]
    if not rows or rows[0][0] != "md-table-excel":
        raise WorkbookError("md-table-excel で書き出したブックではありません")
    data.format_version = rows[0][1]
    data.tool_version = rows[0][2] if len(rows[0]) > 2 else None
    if data.format_version != FORMAT_VERSION:
        raise WorkbookError("このブックの形式（%s）は読めません" % data.format_version)
    for values in rows[1:]:
        if values and values[0] == "markdown":
            data.markdown_path = values[1]
            data.markdown_relpath = values[2] if len(values) > 2 and values[2] else None
    header_index = None
    for index, values in enumerate(rows):
        if values and values[0] == "number":
            header_index = index
            break
    if header_index is None:
        raise WorkbookError("ブックの管理用シートが壊れています")
    for values in rows[header_index + 1:]:
        if not values or values[0] is None:
            continue
        table = SheetTable(int(values[0]))
        table.kind = values[1]
        table.start = int(values[2]) - 1
        table.end = int(values[3]) - 1
        table.source_hash = values[4]
        table.fingerprint = values[5]
        table.sheet_title = values[6]
        table.occurrence = int(values[7]) if len(values) > 7 and values[7] is not None else 0
        data.tables.append(table)


def read(path, force=False):
    """ブックを読む。"""
    if not os.path.splitext(path)[1].lower() in (".xlsx", ".xlsm"):
        raise WorkbookError("Excel のブック（.xlsx）を指定してください: %s" % path)
    if not os.path.exists(path):
        raise WorkbookError("ブックがありません: %s" % path)
    openpyxl = _openpyxl()
    from zipfile import BadZipFile
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            workbook = openpyxl.load_workbook(path, rich_text=True)
    except (BadZipFile, KeyError, ValueError) as error:
        raise WorkbookError("ブックを読めません: %s (%s)" % (path, error))
    except PermissionError:
        raise WorkbookError("ブックを読めません（Excel で開いたままではありませんか）: %s" % path)
    if META_SHEET not in workbook.sheetnames:
        raise WorkbookError("md-table-excel で書き出したブックではありません（管理用シートが無い）")
    data = WorkbookData()
    _read_meta(workbook[META_SHEET], data)
    known_titles = set()
    for table in data.tables:
        name = NAME_PREFIX + str(table.number)
        defined = workbook.defined_names.get(name)
        sheet = None
        start = (1, 1, 1, 1)
        if defined is not None:
            destinations = list(defined.destinations) if defined.type == "RANGE" else []
            for sheet_name, cells in destinations:
                sheet_name = sheet_name.replace("''", "'")
                if sheet_name in workbook.sheetnames:
                    sheet = workbook[sheet_name]
                    bounds = _parse_range(cells)
                    if bounds is not None:
                        start = bounds
                    break
        if sheet is None and table.sheet_title in workbook.sheetnames:
            sheet = workbook[table.sheet_title]
        if sheet is None:
            table.missing = "シートが見つからない（消したか、名前の定義 %s が壊れた）" % name
            continue
        known_titles.add(sheet.title)
        table.sheet_title = sheet.title
        table.model = _read_table_sheet(sheet, start, table, force)
    for title in workbook.sheetnames:
        if title not in known_titles and title not in (INDEX_SHEET, META_SHEET):
            data.warnings.append("シート「%s」はどの表とも対応しないので読まない" % title)
    return data
